import streamlit as st
import pandas as pd
import openpyxl
import io

st.set_page_config(page_title="Atualizador de Orçamento SINAPI", layout="centered")

st.title("🔄 Atualizador Automático de Orçamento")
st.write("Faça o upload da sua Planilha Orçamentária e da Referência do Orçafascio para atualizar os valores automaticamente.")

col1, col2 = st.columns(2)
with col1:
    budget_file = st.file_uploader("📥 1. Planilha Orçamentária Base (.xlsx)", type=["xlsx"])
with col2:
    ref_file = st.file_uploader("📥 2. Referência SINAPI Orçafascio (.xlsx)", type=["xlsx"])

if budget_file and ref_file:
    # Carrega o arquivo para ler o nome das abas antes de processar
    wb = openpyxl.load_workbook(budget_file)
    sheet_names = wb.sheetnames
    
    # Cria uma caixa de seleção para o usuário escolher a aba correta
    selected_sheet = st.selectbox("📌 Selecione qual aba contém o Orçamento a ser atualizado:", sheet_names)

    if st.button("🚀 Processar Atualização"):
        with st.spinner("Analisando e atualizando valores..."):
            try:
                # Carregar o arquivo de referência (Orçafascio) para extrair os novos preços
                ref_df_raw = pd.read_excel(ref_file)
                ref_header_idx = ref_df_raw[ref_df_raw.apply(lambda r: r.astype(str).str.contains("Código", na=False).any(), axis=1)].index[0]
                
                # Ler a referência com o cabeçalho correto
                ref_df = pd.read_excel(ref_file, header=ref_header_idx + 1)
                
                # Renomear colunas baseadas em posição caso os nomes tenham espaços ou variações
                col_codigo_ref = ref_df.columns[1]
                col_valor_ref = ref_df.columns[8]
                
                # Criar um dicionário de preços {codigo: valor}
                ref_df_clean = ref_df.dropna(subset=[col_codigo_ref, col_valor_ref])
                precos_dict = {str(k).strip(): v for k, v in zip(ref_df_clean[col_codigo_ref], ref_df_clean[col_valor_ref])}

                # Selecionar exatamente a aba que o usuário escolheu na caixinha
                ws = wb[selected_sheet]

                # Encontrar onde começam os dados na planilha orçamentária
                header_row = None
                codigo_col = None
                custo_unit_col = None
                bdi_col = None
                preco_unit_col = None
                quantidade_col = None
                preco_total_col = None

                for row in range(1, 30): 
                    for col in range(1, 15): 
                        cell_val = str(ws.cell(row=row, column=col).value).strip().upper()
                        if cell_val == "CÓDIGO":
                            header_row = row
                            codigo_col = col
                        elif cell_val == "CUSTO UNITÁRIO (SEM BDI)":
                            custo_unit_col = col
                        elif cell_val == "VALOR UNITÁRIO (R$)":
                            custo_unit_col = col 
                        elif cell_val == "BDI":
                            bdi_col = col
                        elif cell_val == "PREÇO UNITÁRIO (COM BDI)":
                            preco_unit_col = col
                        elif cell_val == "QUANTIDADE":
                            quantidade_col = col
                        elif cell_val == "PREÇO TOTAL":
                            preco_total_col = col

                if not header_row or not codigo_col:
                    st.error(f"Não foi possível encontrar a coluna 'CÓDIGO' na aba '{selected_sheet}'. Verifique se escolheu a aba correta.")
                    st.stop()
                
                if not custo_unit_col:
                     for col in range(1, 15):
                         if "VALOR UNITÁRIO" in str(ws.cell(row=header_row, column=col).value).upper():
                             custo_unit_col = col
                
                itens_atualizados = 0
                log_alteracoes = [] # Lista para guardar o histórico de modificações

                # Iterar sobre as linhas da planilha de orçamento e atualizar
                for row in range(header_row + 1, ws.max_row + 1):
                    cod_cell = ws.cell(row=row, column=codigo_col)
                    cod_val = str(cod_cell.value).strip() if cod_cell.value else None
                    
                    if cod_val and cod_val in precos_dict:
                        novo_valor = precos_dict[cod_val]
                        valor_antigo = 0.0
                        
                        if custo_unit_col:
                            # Salva o valor antigo antes de sobrescrever
                            celula_antiga = ws.cell(row=row, column=custo_unit_col).value
                            if celula_antiga is not None:
                                try:
                                    valor_antigo = float(celula_antiga)
                                except:
                                    pass
                            
                            # Atualiza para o novo valor
                            ws.cell(row=row, column=custo_unit_col).value = novo_valor
                        
                        bdi_val = 0
                        if bdi_col:
                            bdi_raw = ws.cell(row=row, column=bdi_col).value
                            if bdi_raw is not None:
                                try:
                                    bdi_val = float(bdi_raw) 
                                except:
                                    pass
                        
                        novo_preco_com_bdi = novo_valor * (1 + bdi_val)
                        
                        if preco_unit_col:
                            ws.cell(row=row, column=preco_unit_col).value = novo_preco_com_bdi
                            
                        if quantidade_col and preco_total_col:
                            qtd = ws.cell(row=row, column=quantidade_col).value
                            if qtd is not None:
                                try:
                                    qtd_float = float(qtd)
                                    ws.cell(row=row, column=preco_total_col).value = qtd_float * novo_preco_com_bdi
                                except:
                                    pass
                        
                        # Adiciona o registro no relatório de conferência
                        log_alteracoes.append({
                            "Linha Excel": row,
                            "Código": cod_val,
                            "Valor Antigo (R$)": valor_antigo,
                            "Novo Valor (R$)": novo_valor
                        })
                        
                        itens_atualizados += 1

                # Salvar em memória
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                
                st.success(f"✅ Atualização concluída! {itens_atualizados} itens foram atualizados na aba '{selected_sheet}'.")
                
                # Exibir a tabela de conferência se houver atualizações
                if log_alteracoes:
                    st.write("### 📊 Relatório de Conferência")
                    st.write("Verifique abaixo os valores antigos e os novos valores aplicados:")
                    df_log = pd.DataFrame(log_alteracoes)
                    
                    # Formatar a tabela para exibir como moeda
                    st.dataframe(
                        df_log.style.format({
                            "Valor Antigo (R$)": "{:.2f}",
                            "Novo Valor (R$)": "{:.2f}"
                        }),
                        use_container_width=True
                    )
                
                st.download_button(
                    label="⬇️ Baixar Planilha Atualizada",
                    data=output,
                    file_name="Planilha_Orcamentaria_Atualizada.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            except Exception as e:
                st.error(f"Ocorreu um erro durante o processamento: {e}")
